"""
tests/unit/test_tabular_streaming.py — Streaming tabular pipeline (memory O(batch)).

Covers:
1. reader: scan_tabular_file (pass A) and iter_sheet_batches, and their equivalence
   with the materialising read_csv / read_xlsx / read_tabular_file.
2. state: load_sheets keeps no data rows in the LangGraph state.
3. neo4j_ops: compute_row_normalization_keys (single source of keys for the scan and
   write passes), streaming write path (batches of 500), normalization path
   (passes B/C), bounded normalization tasks, and a live-memory bound.

Everything is faked: no Neo4j, no LLM, no MongoDB.
"""
from __future__ import annotations

import asyncio
import csv
import logging
import tracemalloc
from pathlib import Path
from types import SimpleNamespace

import pytest
from pydantic import BaseModel, Field

from scinr.newton.tabular import neo4j_ops
from scinr.newton.tabular.models import ColumnFieldMapping, ColumnMapping
from scinr.newton.tabular.normalization.engine import NormalizationEngine
from scinr.newton.tabular.reader import (
    iter_sheet_batches,
    preview_indices,
    read_csv,
    read_tabular_file,
    scan_tabular_file,
    select_preview_rows,
)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _write_csv(path: Path, n_rows: int, ncols: int = 3, delimiter: str = ",") -> Path:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=delimiter)
        w.writerow([f"col{c}" for c in range(ncols)])
        for r in range(n_rows):
            w.writerow([f"v{r}_{c}" for c in range(ncols)])
    return path


def _flatten(batches) -> list[list[str]]:
    return [row for batch in batches for row in batch]


def _make_xlsx(path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "People"
    ws1.append(["name", "age"])
    for i in range(12):
        ws1.append([f"p{i}", i])
    ws1.append([None, None])  # empty row in the middle of nowhere
    ws1.append(["last", 99])
    ws2 = wb.create_sheet("Empty")  # no cells at all
    ws3 = wb.create_sheet("Tiny")
    ws3.append(["a", "b", "c"])
    ws3.append([1, 2])  # short row -> padded
    del ws2
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# 1. Reader
# ---------------------------------------------------------------------------


class TestPreviewIndices:
    @pytest.mark.parametrize("n", [0, 1, 4, 5, 6, 7, 8, 100, 1001])
    def test_matches_select_preview_rows(self, n):
        rows = [[str(i)] for i in range(n)]
        sheet = {"sheet_name": "s", "headers": ["h"], "all_rows": rows, "total_rows": n}

        assert select_preview_rows(sheet)["row_indices"] == preview_indices(n)


class TestScanTabularFile:
    def test_csv_scan_has_headers_total_and_preview_without_rows(self, tmp_path):
        p = _write_csv(tmp_path / "data.csv", 1000)

        (scan,) = scan_tabular_file(p)

        assert scan["sheet_name"] == "data"
        assert scan["headers"] == ["col0", "col1", "col2"]
        assert scan["total_rows"] == 1000
        assert "all_rows" not in scan
        preview = scan["preview"]
        assert preview["row_indices"] == [0, 250, 500, 750, 999]
        assert [r[0] for r in preview["preview_rows"]] == ["v0_0", "v250_0", "v500_0", "v750_0", "v999_0"]
        assert preview["total_rows"] == 1000

    @pytest.mark.parametrize("n", [1, 3, 5])
    def test_small_csv_previews_every_row(self, tmp_path, n):
        p = _write_csv(tmp_path / "small.csv", n)

        (scan,) = scan_tabular_file(p)

        assert scan["total_rows"] == n
        assert scan["preview"]["row_indices"] == list(range(n))
        assert len(scan["preview"]["preview_rows"]) == n

    def test_total_ignores_empty_rows_and_counts_multiline_records_once(self, tmp_path):
        p = tmp_path / "odd.csv"
        p.write_text('id,note\n1,"a\nb"\n\n,\n2,c\n', encoding="utf-8")

        (scan,) = scan_tabular_file(p)

        assert scan["total_rows"] == 2
        assert scan["preview"]["preview_rows"] == [["1", "a\nb"], ["2", "c"]]

    def test_empty_csv_gives_placeholder_sheet(self, tmp_path):
        p = tmp_path / "empty.csv"
        p.write_text("", encoding="utf-8")

        (scan,) = scan_tabular_file(p)

        assert scan["sheet_name"] == "empty"
        assert scan["headers"] == []
        assert scan["total_rows"] == 0
        assert scan["preview"]["preview_rows"] == []

    def test_scan_agrees_with_read_tabular_file_for_xlsx(self, tmp_path):
        p = _make_xlsx(tmp_path / "book.xlsx")

        scans = scan_tabular_file(p)
        sheets = read_tabular_file(p)

        assert [s["sheet_name"] for s in scans] == [s["sheet_name"] for s in sheets] == ["People", "Tiny"]
        for scan, sheet in zip(scans, sheets, strict=True):
            assert scan["headers"] == sheet["headers"]
            assert scan["total_rows"] == sheet["total_rows"]
            assert scan["preview"] == select_preview_rows(sheet)

    def test_unsupported_extension_raises(self, tmp_path):
        p = tmp_path / "x.txt"
        p.write_text("a,b")
        with pytest.raises(ValueError):
            scan_tabular_file(p)


class TestIterSheetBatches:
    def test_concatenation_equals_read_csv_for_small_and_large_files(self, tmp_path):
        for n in (1, 7, 501, 1234):
            p = _write_csv(tmp_path / f"f{n}.csv", n, ncols=4)
            expected = read_csv(p)[0]["all_rows"]

            got = _flatten(iter_sheet_batches(p, p.stem, batch_size=500))

            assert got == expected

    def test_batches_are_bounded_ordered_and_last_one_partial(self, tmp_path):
        p = _write_csv(tmp_path / "f.csv", 1201)

        batches = list(iter_sheet_batches(p, "f", batch_size=500))

        assert [len(b) for b in batches] == [500, 500, 201]
        assert batches[0][0][0] == "v0_0"
        assert batches[2][-1][0] == "v1200_0"

    def test_header_only_and_empty_sheets_yield_nothing(self, tmp_path):
        only_header = tmp_path / "h.csv"
        only_header.write_text("a,b\n", encoding="utf-8")
        empty = tmp_path / "e.csv"
        empty.write_text("", encoding="utf-8")

        assert list(iter_sheet_batches(only_header, "h")) == []
        assert list(iter_sheet_batches(empty, "e")) == []

    def test_short_and_long_rows_are_padded_and_trimmed_like_read_csv(self, tmp_path):
        p = tmp_path / "ragged.csv"
        p.write_text("a,b,c\n1\n1,2,3,4,5\n\n7,8,9\n", encoding="utf-8")

        assert _flatten(iter_sheet_batches(p, "ragged")) == read_csv(p)[0]["all_rows"]
        assert read_csv(p)[0]["all_rows"] == [["1", "", ""], ["1", "2", "3"], ["7", "8", "9"]]

    def test_xlsx_sheets_stream_the_same_rows_as_read_tabular_file(self, tmp_path):
        p = _make_xlsx(tmp_path / "book.xlsx")
        by_name = {s["sheet_name"]: s["all_rows"] for s in read_tabular_file(p)}

        for name in ("People", "Tiny"):
            assert _flatten(iter_sheet_batches(p, name, batch_size=5)) == by_name[name]

    def test_unknown_sheet_raises_key_error(self, tmp_path):
        p = _write_csv(tmp_path / "f.csv", 3)
        with pytest.raises(KeyError):
            list(iter_sheet_batches(p, "nope"))

    def test_file_is_closed_when_the_iterator_is_abandoned(self, tmp_path):
        p = _write_csv(tmp_path / "f.csv", 2000)
        it = iter_sheet_batches(p, "f", batch_size=10)
        next(it)
        it.close()  # must not raise, and must release the file handle


# ---------------------------------------------------------------------------
# 2. LangGraph state carries no rows
# ---------------------------------------------------------------------------


class TestLoadSheetsKeepsNoRows:
    async def test_state_sheets_have_metadata_only(self, tmp_path, monkeypatch):
        from scinr.newton.tabular.nodes import load_sheets

        async def _noop(*a, **k):
            return None

        monkeypatch.setattr("scinr.newton.annotation.neo4j_ops.ensure_catalog_models_once", _noop)
        monkeypatch.setattr("scinr.newton.annotation.neo4j_ops.ensure_theme_structure_once", _noop)
        monkeypatch.setattr("scinr.newton.ingest.config.get_async_driver", lambda: object())
        monkeypatch.setattr("scinr.newton.utils.theme_registry.get_theme_registry", lambda: object())

        p = _write_csv(tmp_path / "big.csv", 3000)
        state = {"file_path": str(p), "document_name": "big", "doc_path": "big", "errors": []}

        out = await load_sheets(state)

        (sheet,) = out["sheets"]
        assert "all_rows" not in sheet
        assert sheet["file_path"] == str(p)
        assert sheet["sheet_name"] == "big"
        assert sheet["total_rows"] == 3000
        assert sheet["headers"] == ["col0", "col1", "col2"]
        assert "| v0_0 |" in sheet["preview_markdown"]
        assert "showing 5 of 3000 rows" in sheet["preview_markdown"]


# ---------------------------------------------------------------------------
# 3. neo4j_ops — streaming writes
# ---------------------------------------------------------------------------


class NormalizedAddress(BaseModel):
    city: str | None = None


class Person(BaseModel):
    name: str | None = None
    address_raw: str | None = None
    address: NormalizedAddress | None = Field(
        default=None,
        json_schema_extra={
            "normalization_model": True,
            "normalization_source_fields": ["address_raw"],
        },
    )


class PersonComposite(BaseModel):
    person: Person | None = None


_MAPPING = ColumnMapping(
    mappings=[
        ColumnFieldMapping(column_name="name", model_field_name="name", target_model="primary"),
        ColumnFieldMapping(
            column_name="addr", model_field_name="address_raw", target_model="primary"
        ),
    ]
)


def _factory(rows: list[list[str]], counter: list[int] | None = None):
    """RowBatchFactory over an in-memory list; counts how many passes were started."""

    def row_batches(batch_size: int):
        if counter is not None:
            counter[0] += 1
        for i in range(0, len(rows), batch_size):
            yield rows[i : i + batch_size]

    return row_batches


class _FakeResult:
    async def single(self):
        # write_tabular_subgraph checks the owning :Document was matched.
        return {"created": 1}


class _FakeTx:
    async def run(self, *a, **k):
        return _FakeResult()

    async def commit(self):
        return None

    async def rollback(self):
        return None


class _FakeSession:
    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False

    async def begin_transaction(self):
        return _FakeTx()

    async def run(self, *a, **k):
        return None


class _FakeDriver:
    def session(self, **kwargs):
        return _FakeSession()


@pytest.fixture
def fake_cfg(monkeypatch):
    cfg = SimpleNamespace(
        neo4j_database="neo4j",
        normalization_enabled=True,
        normalization_llm=object(),
        llm=None,
        normalization_batch_size=5,
        llm_concurrency=2,
    )
    monkeypatch.setattr(neo4j_ops, "get_config", lambda: cfg)
    monkeypatch.setattr("scinr.newton.config.get_llm_semaphore", lambda: asyncio.Semaphore(2))
    return cfg


@pytest.fixture
def recorded_batches(monkeypatch):
    """Replace _write_row_batch; record (batch_len, batch_start_index, composite_results)."""
    calls: list[dict] = []

    async def fake_write_row_batch(**kwargs):
        calls.append(
            {
                "n": len(kwargs["rows_batch"]),
                "start": kwargs["batch_start_index"],
                "rows": kwargs["rows_batch"],
                "composites": kwargs.get("composite_results"),
                "row_indices": kwargs.get("row_indices"),
            }
        )

    monkeypatch.setattr(neo4j_ops, "_write_row_batch", fake_write_row_batch)
    return calls


class TestComputeRowNormalizationKeys:
    def test_key_matches_the_engine_hash_and_is_deterministic(self):
        keys = neo4j_ops.compute_row_normalization_keys(
            ["name", "addr"], ["Ann", "1 Main St"], _MAPPING, Person, {}, []
        )
        again = neo4j_ops.compute_row_normalization_keys(
            ["name", "addr"], ["Ann", "1 Main St"], _MAPPING, Person, {}, []
        )

        (key,) = keys
        assert key.unique_key == (
            "NormalizedAddress:"
            + NormalizationEngine._hash_source_values({"address_raw": "1 Main St"})
        )
        assert key.is_primary and key.field_name == "address"
        assert key.target_type is NormalizedAddress
        assert key.source_values == {"address_raw": "1 Main St"}
        assert keys == again

    def test_row_without_normalizable_value_has_no_keys(self):
        assert neo4j_ops.compute_row_normalization_keys(
            ["name", "addr"], ["Ann", ""], _MAPPING, Person, {}, []
        ) == []

    def test_scan_pass_and_write_pass_use_the_same_keys(self):
        headers = ["name", "addr"]
        rows = [["a", "x"], ["b", "y"], ["c", "x"], ["d", ""]]

        dedup_map, scanned = neo4j_ops._build_normalization_dedup_map(
            headers, [rows[:2], rows[2:]], Person, {}, [], _MAPPING
        )
        write_pass_keys = {
            k.unique_key
            for r in rows
            for k in neo4j_ops.compute_row_normalization_keys(headers, r, _MAPPING, Person, {}, [])
        }

        assert scanned == 4
        assert set(dedup_map) == write_pass_keys
        assert len(dedup_map) == 2  # "x" is deduplicated across rows 0 and 2
        entry = dedup_map[next(iter(dedup_map))]
        assert entry.row_indices == []  # no per-row bookkeeping any more


class TestStandardWritePathStreams:
    async def test_2000_rows_are_written_in_4_batches_of_500_never_the_whole_list(
        self, fake_cfg, recorded_batches, monkeypatch
    ):
        async def fake_write_annotation(*a, **k):
            return None

        monkeypatch.setattr(
            "scinr.newton.annotation.neo4j_ops.write_annotation", fake_write_annotation
        )
        rows = [[f"n{i}", f"a{i}"] for i in range(2000)]
        passes = [0]
        decision = SimpleNamespace(
            matched_model_class=None, complementary_models=[], supplementary_fields=[]
        )
        sheet = {
            "file_path": "unused",
            "sheet_name": "s",
            "headers": ["name", "addr"],
            "total_rows": 2000,
        }

        table_id = await neo4j_ops.write_tabular_subgraph(
            driver=_FakeDriver(),
            doc_path="doc",
            document_name="doc",
            resolved_version=1,
            sheet=sheet,
            row_batches=_factory(rows, passes),
            sheet_index=0,
            decision=decision,
            mapping=_MAPPING,
            tenant_id="acme",
        )

        assert table_id == "acme::doc::1::table_1"
        assert [c["n"] for c in recorded_batches] == [500, 500, 500, 500]
        assert [c["start"] for c in recorded_batches] == [0, 500, 1000, 1500]
        assert all(c["n"] < len(rows) for c in recorded_batches)
        assert passes == [1]  # a single streaming pass when there is no normalization


class TestNormalizationWritePath:
    @pytest.fixture
    def engine_calls(self, monkeypatch):
        calls: list[list[str]] = []

        async def fake_process_key_batch(self, entries):
            calls.append([e.unique_key for e in entries])
            for e in entries:
                self.result_cache[e.unique_key] = NormalizedAddress(
                    city=str(e.source_values["address_raw"]).upper()
                )
            return list(entries)

        monkeypatch.setattr(NormalizationEngine, "process_key_batch", fake_process_key_batch)
        return calls

    async def _run(self, rows, passes):
        return await neo4j_ops._write_tabular_with_normalization(
            driver=_FakeDriver(),
            table_composite_id="doc::1::table_1",
            headers=["name", "addr"],
            row_batches=_factory(rows, passes),
            decision_uid="uid",
            decision=SimpleNamespace(matched_model_class="Person"),
            mapping=_MAPPING,
            primary_cls=Person,
            composite_cls=PersonComposite,
            primary_field_name="person",
            comp_class_names=[],
            comp_cls_map={},
            document_name="doc",
            theme="default",
            sheet_page_id="",
        )

    async def test_scan_normalize_write_passes_apply_the_cached_results(
        self, fake_cfg, recorded_batches, engine_calls, caplog
    ):
        # 1200 rows, 7 distinct addresses, every 4th row without address.
        cities = [f"city {i}" for i in range(7)]
        rows = [
            [f"n{i}", "" if i % 4 == 3 else cities[i % 7]] for i in range(1200)
        ]
        passes = [0]

        with caplog.at_level(logging.WARNING, logger=neo4j_ops.logger.name):
            written = await self._run(rows, passes)

        assert written == 1200
        assert passes == [2]  # pass B (scan) + pass C (write)
        # LLM sees only the unique keys, in batches of normalization_batch_size=5.
        seen_keys = [k for batch in engine_calls for k in batch]
        assert len(seen_keys) == len(set(seen_keys)) == 7
        assert all(len(batch) <= 5 for batch in engine_calls)
        # Batches of 500 in file order, indices contiguous.
        assert [c["n"] for c in recorded_batches] == [500, 500, 200]
        assert [c["start"] for c in recorded_batches] == [0, 500, 1000]
        assert recorded_batches[1]["row_indices"] == list(range(500, 1000))
        # Normalization applied where there is a key; plain otherwise.
        for c in recorded_batches:
            for offset, row_values, composite, extraction_uid in c["composites"]:
                row_index = c["start"] + offset
                assert row_values == rows[row_index]
                assert extraction_uid  # deterministic uid
                if row_index % 4 == 3:
                    assert composite.person.address is None
                else:
                    assert composite.person.address == NormalizedAddress(
                        city=cities[row_index % 7].upper()
                    )
        assert not [r for r in caplog.records if "cache miss" in r.getMessage()]

    async def test_llm_failure_leaves_the_field_as_is_and_still_writes_every_row(
        self, fake_cfg, recorded_batches, monkeypatch, caplog
    ):
        async def failing_process_key_batch(self, entries):
            return []  # nothing cached: normalization "failed"

        monkeypatch.setattr(NormalizationEngine, "process_key_batch", failing_process_key_batch)
        rows = [[f"n{i}", "addr"] for i in range(10)]

        with caplog.at_level(logging.WARNING, logger=neo4j_ops.logger.name):
            written = await self._run(rows, [0])

        assert written == 10
        (call,) = recorded_batches
        assert all(c[2].person.address is None for c in call["composites"])
        assert len([r for r in caplog.records if "cache miss" in r.getMessage()]) == 10

    async def test_no_normalizable_values_falls_back_to_the_standard_path(
        self, fake_cfg, recorded_batches, engine_calls
    ):
        rows = [[f"n{i}", ""] for i in range(700)]
        passes = [0]

        written = await self._run(rows, passes)

        assert written == 700
        assert engine_calls == []
        assert passes == [2]  # scan pass + standard write pass
        assert [c["n"] for c in recorded_batches] == [500, 200]
        assert all(c["composites"] is None for c in recorded_batches)


class TestRunBounded:
    async def test_at_most_limit_tasks_alive_and_all_items_processed(self):
        alive = 0
        max_alive = 0
        done: list[int] = []

        async def worker(i: int) -> None:
            nonlocal alive, max_alive
            alive += 1
            max_alive = max(max_alive, alive)
            await asyncio.sleep(0.001)
            alive -= 1
            done.append(i)

        await neo4j_ops._run_bounded(range(200), worker, limit=8)

        assert max_alive <= 8
        assert sorted(done) == list(range(200))

    async def test_items_are_consumed_lazily(self):
        consumed = 0

        def items():
            nonlocal consumed
            for i in range(100):
                consumed += 1
                yield i

        started = asyncio.Event()
        release = asyncio.Event()

        async def worker(i: int) -> None:
            started.set()
            await release.wait()

        task = asyncio.create_task(neo4j_ops._run_bounded(items(), worker, limit=3))
        await started.wait()
        await asyncio.sleep(0.01)
        assert consumed <= 4  # 3 running + the one waiting for a free slot
        release.set()
        await task
        assert consumed == 100

    async def test_first_failure_is_raised_and_pending_tasks_are_cancelled(self):
        cancelled = 0

        async def worker(i: int) -> None:
            nonlocal cancelled
            if i == 0:
                raise RuntimeError("boom")
            try:
                await asyncio.sleep(10)
            except asyncio.CancelledError:
                cancelled += 1
                raise

        # If the failure were only surfaced once every task finished, this would take 10 s.
        with pytest.raises(RuntimeError, match="boom"):
            await asyncio.wait_for(neo4j_ops._run_bounded(range(3), worker, limit=3), timeout=2)
        assert cancelled == 2

    async def test_empty_input_is_a_noop(self):
        async def worker(i: int) -> None:  # pragma: no cover
            raise AssertionError

        await neo4j_ops._run_bounded([], worker, limit=4)


class TestLiveMemoryIsIndependentOfRowCount:
    async def test_standard_write_path_peak_is_bounded(self, tmp_path, fake_cfg, monkeypatch):
        n_rows = 150_000
        p = _write_csv(tmp_path / "huge.csv", n_rows, ncols=8)

        async def sink(**kwargs):
            return None  # like a real write: the batch is dropped afterwards

        monkeypatch.setattr(neo4j_ops, "_write_row_batch", sink)

        def factory(batch_size):
            return iter_sheet_batches(p, "huge", batch_size)

        tracemalloc.start()
        try:
            tracemalloc.reset_peak()
            written = await neo4j_ops._write_rows_in_batches(
                driver=None,
                table_composite_id="t",
                headers=["c"],
                row_batches=factory,
                decision_uid="u",
                decision=None,
                mapping=None,
                primary_cls=None,
                composite_cls=None,
                primary_field_name=None,
                comp_class_names=[],
                comp_cls_map={},
                document_name="d",
                theme="default",
                sheet_page_id="",
            )
            (scan,) = scan_tabular_file(p)
            peak_mb = tracemalloc.get_traced_memory()[1] / 1024 / 1024
        finally:
            tracemalloc.stop()

        assert written == n_rows
        assert scan["total_rows"] == n_rows
        # Materialising 150k x 8 cells costs well over 100 MB; streaming stays tiny.
        assert peak_mb < 50, f"peak live memory {peak_mb:.0f} MB"
